#!/usr/bin/env python
"""Score a completed blind annotation sheet against the Exp 1 gold labels (AES 2026).

Reports Cohen's kappa and raw agreement overall and separately for aligned and
adversarial cases, plus a list of disagreements for adjudication.

    python scripts/compute_kappa.py routing_annotation_sheet.xlsx
"""

from __future__ import annotations

import argparse
import sys
from collections import Counter
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "scripts"))
sys.path.insert(0, str(ROOT / "src"))

from make_annotation_sheet import blind_order  # noqa: E402


def cohen_kappa(a: list, b: list) -> float:
    n = len(a)
    po = sum(x == y for x, y in zip(a, b)) / n
    ca, cb = Counter(a), Counter(b)
    pe = sum(ca[k] * cb[k] for k in set(ca) | set(cb)) / (n * n)
    return 1.0 if pe == 1 else (po - pe) / (1 - pe)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("sheet")
    args = ap.parse_args()

    from openpyxl import load_workbook

    ws = load_workbook(args.sheet, data_only=True)["Label"]
    labels = {int(r[0]): (str(r[2]).strip().lower() if r[2] else "")
              for r in ws.iter_rows(min_row=2, values_only=True) if r[0] is not None}
    order = blind_order()
    missing = [i for i in range(1, len(order) + 1) if not labels.get(i)]
    if missing:
        raise SystemExit(f"unlabelled items: {missing}")

    rows = []
    for item, (_, (task, gold)) in enumerate(order, start=1):
        rows.append((item, task.metadata["case"], gold.value, labels[item], task.description))

    def report(name: str, subset: list) -> None:
        g = [r[2] for r in subset]
        a = [r[3] for r in subset]
        agree = sum(x == y for x, y in zip(g, a))
        print(f"{name:12s} n={len(subset):2d}  agreement={agree}/{len(subset)} "
              f"({agree / len(subset):.2f})  kappa={cohen_kappa(g, a):.3f}")

    report("overall", rows)
    report("aligned", [r for r in rows if r[1] == "aligned"])
    report("adversarial", [r for r in rows if r[1] == "adversarial"])
    dis = [r for r in rows if r[2] != r[3]]
    if dis:
        print("\nDisagreements (item, case, gold, annotator, description):")
        for r in dis:
            print(f"  #{r[0]:2d} {r[1]:11s} gold={r[2]:15s} ann={r[3]:15s} {r[4][:80]}")


if __name__ == "__main__":
    main()
