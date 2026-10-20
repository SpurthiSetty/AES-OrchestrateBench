#!/usr/bin/env python
"""Build a blind second-annotator sheet for the Exp 1 routing diagnostic (AES 2026).

The sheet shows each task description and the four route definitions, in a fixed
shuffled order, with no gold label, no aligned/adversarial tag, and no metadata
flags (the gold standard is the description's intent, and the flags are
deliberately misleading on adversarial cases). The annotator picks one route per
row from a dropdown. Score it with ``scripts/compute_kappa.py``.

    python scripts/make_annotation_sheet.py --out routing_annotation_sheet.xlsx
"""

from __future__ import annotations

import argparse
import random
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "src"))

from orchestratebench.routing_comparison import _DECISION_GUIDE, make_routing_eval_set  # noqa: E402

SHUFFLE_SEED = 2026


def blind_order() -> list:
    cases = list(enumerate(make_routing_eval_set()))
    random.Random(SHUFFLE_SEED).shuffle(cases)
    return cases


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--out", default="routing_annotation_sheet.xlsx")
    args = ap.parse_args()

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = Workbook()
    ws = wb.active
    ws.title = "Label"
    routes = [d.value for d in _DECISION_GUIDE]
    ws.append(["item", "task description", "your route", "notes (optional)"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for item, (_, (task, _gold)) in enumerate(blind_order(), start=1):
        ws.append([item, task.description, "", ""])
    dv = DataValidation(type="list", formula1='"' + ",".join(routes) + '"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"C2:C{ws.max_row}")
    fill = PatternFill("solid", fgColor="FFF4CC")
    for r in range(2, ws.max_row + 1):
        ws[f"C{r}"].fill = fill
        ws[f"B{r}"].alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 90
    ws.column_dimensions["C"].width = 18
    ws.column_dimensions["D"].width = 30
    ws.freeze_panes = "A2"

    guide = wb.create_sheet("Instructions")
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 100
    lines = [
        ("Task", "For each row on the Label sheet, choose the ONE route that best fits what the "
                 "task actually requires, judging only from the description."),
        ("Blind", "Work independently. Do not look at the gold labels, the repo's eval set, or "
                  "anyone else's labels before finishing."),
        ("", ""),
        ("Routes", ""),
    ] + [(d.value, desc) for d, desc in _DECISION_GUIDE.items()]
    for a, b in lines:
        guide.append([a, b])
        guide[f"B{guide.max_row}"].alignment = Alignment(wrap_text=True)
    for r in (1, 2, 4):
        guide[f"A{r}"].font = Font(bold=True)
    wb.move_sheet("Instructions", offset=-1)
    wb.save(args.out)
    print(f"wrote {args.out} ({ws.max_row - 1} items)")


if __name__ == "__main__":
    main()
